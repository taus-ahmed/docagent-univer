"""
Regression tests for defects that have been fixed. They must stay green.

── P3: fabricated duplicate line item on the structural path ─────────────────

Audit 2026-08-17 §6.4, reproduced live 4/4 on BalanceSheet-synth.pdf under the
then-production config. The exported sheet contained a "Cash & Cash
Equivalents / 168000" row twice; the document lists it once. Mechanism:

  1. Layer 1 emitted phantom "Total Current Assets" sections (item_count=1)
     alongside the real ones.
  2. Layer 2, asked to extract the phantom, returned the whole column (5 rows).
  3. S5 truncation kept rows[:1] — the section's FIRST LINE ITEM, not its total.
  4. _match_group's substring test routed "Total Current Assets" into the real
     "Current Assets" group and appended the row.
  5. Neither dedup guard fired: content-signature needed identical content
     (1 row != 4); superset dedup needed len(r1) > len(r2) and the truncated
     phantom was a SUBSET.

STATUS: fixed in Phase 2. Every one of those five components has been deleted
along with the layout path (2c/2d). The pipeline that replaced it cannot
reach the same state — a row is written only to the address it was requested
for, and a fabricated row is caught by its source span rather than by
comparing sections to each other.

This test therefore no longer replays the old Layer-1/Layer-2 responses (the
code that consumed them is gone). It asserts the SAME THING the original
asserted — the export must not contain a line item the document does not have
— against the current pipeline, and it feeds the model response that used to
produce the phantom: a duplicate row claiming a document line another row
already used.
"""
import contextlib
import io
import json

import pytest

from tests.harness import bootstrap as bs

bs.bootstrap()

FIXTURE_PDF = bs.BACKEND_DIR / "data" / "test_uploads" / "BalanceSheet-synth.pdf"
FIXTURE_TEMPLATE = bs.BACKEND_DIR / "data" / "test_uploads" / "bs_unlabeled_template.json"

# Lines exactly as pdfplumber reads them out of the fixture.
LINES = {
    "cash": "Cash & Cash Equivalents 168,000 Accounts Payable 262,000",
    "ar": "Accounts Receivable 95,000 Short-Term Debt 80,000",
    "inv": "Inventory 140,000 Accrued Expenses 33,000",
    "prepaid": "Prepaid Expenses 12,000 Deferred Revenue 25,000",
    "total": "Total Current Assets 415,000 Total Current Liabilities 400,000",
}


def _row(line, assets_label, assets_amt, liab_label, liab_amt):
    """A table row in the shape the slot prompt asks for. The band's two value
    columns are both headed "Amount", so they are addressed by column letter."""
    return {"cells": {"Current Assets": assets_label, "Amount (B)": assets_amt,
                      "Current Liabilities": liab_label, "Amount (D)": liab_amt},
            "source": line, "page": 1}


REAL_ROWS = [
    _row(LINES["cash"], "Cash & Cash Equivalents", "168,000", "Accounts Payable", "262,000"),
    _row(LINES["ar"], "Accounts Receivable", "95,000", "Short-Term Debt", "80,000"),
    _row(LINES["inv"], "Inventory", "140,000", "Accrued Expenses", "33,000"),
    _row(LINES["prepaid"], "Prepaid Expenses", "12,000", "Deferred Revenue", "25,000"),
]

# The fabrication: a fifth row repeating the first document line. This is what
# the old pipeline produced after S5 truncated a phantom "total" section down
# to that section's first line item.
PHANTOM_ROW = _row(LINES["cash"], "Cash & Cash Equivalents", "168,000",
                   "Accounts Payable", "262,000")

RESPONSE = {
    "fields": {
        "F1": {"value": "415,000", "source": LINES["total"], "page": 1},
        "F2": {"value": "400,000", "source": LINES["total"], "page": 1},
    },
    "tables": {"table": REAL_ROWS + [PHANTOM_ROW]},
}


@pytest.fixture()
def model_returns_a_phantom_row(monkeypatch):
    from connectors.groq_client import LLMResponse
    from connectors.llm_router import LLMRouter

    def extract(self, text="", image_b64="", prompt="", system_instruction="",
                model=None):
        return LLMResponse(raw_text=json.dumps(RESPONSE), parsed_json=RESPONSE,
                           model_used="stub", tokens_used=0, success=True)

    monkeypatch.setattr(LLMRouter, "extract", extract)
    yield


def _run_and_export():
    """Run the real pipeline on the fixture, then export exactly as the
    download endpoint does — rebuilt from the persisted extraction_json."""
    import openpyxl
    from app.api.routes.extract import (_analyse_template_regions,
                                        _extract_with_template, _parse_template,
                                        _write_excel)
    from app.models.models import ColumnTemplate, DocumentResult
    from orchestrator import Orchestrator
    from app.api.routes.templates import _compute_and_store_shape
    from tests.harness.runner import _schema_path

    bs.chdir_backend()
    grid = json.loads(FIXTURE_TEMPLATE.read_text(encoding="utf-8"))
    tpl = ColumnTemplate(name="bs_unlabeled", document_type="balance_sheet",
                         description=json.dumps(grid), columns_json="[]")
    _compute_and_store_shape(tpl)
    template_data = _parse_template(tpl)

    log = io.StringIO()
    with contextlib.redirect_stdout(log):
        orchestrator = Orchestrator(client_schema_path=_schema_path())
        results = _extract_with_template(orchestrator, FIXTURE_PDF, template_data)

    docs = []
    for r in results:
        ed = getattr(r, "extracted_data", None) or {}
        docs.append(DocumentResult(
            filename=getattr(r, "filename", FIXTURE_PDF.name),
            document_type=getattr(r, "document_type", "") or "balance_sheet",
            extraction_json=json.dumps(ed, default=str)))
    wb = openpyxl.Workbook()
    ws = wb.active
    with contextlib.redirect_stdout(io.StringIO()):
        _write_excel(ws, docs, grid, _analyse_template_regions(grid), openpyxl)
    return results, ws, log.getvalue()


class TestP3FabricatedLineItem:
    def test_no_fabricated_duplicate_row_in_export(self, model_returns_a_phantom_row):
        results, ws, log = _run_and_export()

        # Guards: the intended path must actually have run. Without these the
        # test could pass because nothing was extracted at all — which is how
        # this test caught the layout path's removal rather than silently
        # going green on an empty sheet.
        assert "-> slot path" in log, log
        assert "falling back to legacy" not in log, log
        assert results and getattr(results[0], "success", False), (
            getattr(results[0], "error_message", "no results"))

        rows = []
        for row in ws.iter_rows(min_row=2):          # row 1 = the band header
            label = row[0].value
            if label is not None and str(label).strip():
                rows.append(str(label).strip())

        duplicates = sorted({r for r in rows if rows.count(r) > 1})
        assert duplicates == [], (
            f"Fabricated duplicate line item(s) in the exported sheet: "
            f"{duplicates}. The document lists each asset exactly once; the "
            f"export wrote: {rows}. This is audit finding P3."
        )

    def test_the_phantom_row_is_reported_not_silently_dropped(
            self, model_returns_a_phantom_row):
        """Dropping it quietly would be a different bug. The document must say
        that a row was discarded and why."""
        results, _ws, _log = _run_and_export()
        notes = (results[0].extracted_data or {}).get("validation_notes", [])
        assert any("duplicate row dropped" in n for n in notes), notes

    def test_the_four_real_rows_all_survive(self, model_returns_a_phantom_row):
        """The guard must remove the fabrication WITHOUT removing real data —
        the failure mode the original superset dedup had in the other
        direction."""
        results, ws, _log = _run_and_export()
        labels = [str(r[0].value).strip() for r in ws.iter_rows(min_row=2)
                  if r[0].value and str(r[0].value).strip()]
        for expected in ["Cash & Cash Equivalents", "Accounts Receivable",
                         "Inventory", "Prepaid Expenses"]:
            assert expected in labels, (expected, labels)

    def test_both_side_by_side_totals_are_written(self, model_returns_a_phantom_row):
        """The fixture puts two label/value pairs on one row. Taking only the
        leftmost silently loses half the sheet."""
        _results, ws, _log = _run_and_export()
        values = [c.value for row in ws.iter_rows() for c in row]
        assert 415000 in values, values
        assert 400000 in values, values
