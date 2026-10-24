"""The HTTP worker must preserve the OCR engine's document contract."""

from __future__ import annotations

import io
from types import SimpleNamespace

import openpyxl
import pytest

from tests.harness import bootstrap as bs

bs.bootstrap()


def _job(db, owner):
    from app.models.models import ExtractionJob

    job = ExtractionJob(
        user_id=owner.id,
        client_id="acme_001",
        status="pending",
        total_docs=1,
        input_source="upload",
    )
    db.add(job)
    db.commit()
    db.refresh(job)
    return job


def _result(data, *, success=True, filename="stored.png"):
    return SimpleNamespace(
        filename=filename,
        document_type="cheque",
        success=success,
        extracted_data=data,
        validation=SimpleNamespace(errors=[], needs_review=False),
        extraction_response=None,
        processing_time_ms=1,
        error="" if success else "engine failed",
    )


def _run_worker(monkeypatch, tmp_path, db_url, job_id, engine_results,
                *, template_data=None, legacy_results=None):
    from app.api.routes import extract as ex
    import orchestrator

    image = tmp_path / f"route-{job_id}.png"
    image.write_bytes(b"not read: extraction is stubbed")
    calls = []

    monkeypatch.setattr(orchestrator, "Orchestrator",
                        lambda **_kwargs: SimpleNamespace())
    monkeypatch.setattr("app.services.job_worker._resolve_source", lambda _storage, _key: image)

    def engine(_orchestrator, path, template, **kwargs):
        calls.append(("engine", path, template, kwargs))
        return engine_results

    def legacy(_orchestrator, path, template):
        calls.append(("legacy", path, template, {}))
        return legacy_results or engine_results

    monkeypatch.setattr("app.services.job_worker._extract_with_template", engine)
    monkeypatch.setattr("app.services.job_worker._extract_image_with_template", legacy)
    ex._run_extraction_sync(
        job_id,
        ["stored/key.png"],
        "unused-schema.yaml",
        db_url,
        template_data,
        str(bs.REPO_DIR),
        str(bs.BACKEND_DIR),
        str(bs.ENGINE_DIR),
        original_filenames=["Customer cheque.png"],
    )
    return calls


@pytest.mark.parametrize("template_data", [None, {"doc_type": "cheque"}])
def test_standalone_images_use_the_engine_with_or_without_a_template(
        monkeypatch, tmp_path, _db_url, db, users, template_data):
    job = _job(db, users["acme"])
    calls = _run_worker(
        monkeypatch,
        tmp_path,
        _db_url,
        job.id,
        [_result({"extracted_data": {}, "validation": {}})],
        template_data=template_data,
    )

    assert [call[0] for call in calls] == ["engine"]
    assert calls[0][2] == template_data


def test_an_image_engine_failure_falls_back_to_the_legacy_path(
        monkeypatch, tmp_path, _db_url, db, users):
    job = _job(db, users["acme"])
    calls = _run_worker(
        monkeypatch,
        tmp_path,
        _db_url,
        job.id,
        [_result({}, success=False)],
        legacy_results=[_result({"extracted_data": {}, "validation": {}})],
    )

    assert [call[0] for call in calls] == ["engine", "legacy"]


def test_multi_cheque_metadata_round_trips_and_exports_one_row_per_cheque(
        monkeypatch, tmp_path, _db_url, db, users, client, auth):
    job = _job(db, users["acme"])
    common_validation = {
        "source_region": {"index": 1, "of": 2, "box": [0, 1, 2, 3]},
        "possible_multiple_documents": {"bands": 2},
        "routing_checksum": "pass",
        "ocr_pages": [{
            "page": 1,
            "words": 37,
            "skew_detector": "osd",
            "retry": {"attempted": True, "selected": "deskewed"},
        }],
        "ocr_failed_pages": [],
        "ocr_zero_word_pages": [],
        "page_origins": {"1": "ocr"},
        "page_gates": {"1": {"origin": "ocr", "confidence_ceiling": "unverified"}},
    }

    def cheque(index, number, review):
        validation = dict(common_validation)
        validation["source_region"] = dict(common_validation["source_region"],
                                            index=index)
        return _result({
            "document_index": index,
            "document_count": 2,
            "source_pages": [1, 1],
            "extracted_data": {
                "Cheque Number": {
                    "value": number,
                    "confidence": "unverified",
                    "review_required": review,
                    "review_reasons": ["ocr_origin"] if review else [],
                    "amount_cross_check": "pass",
                    "aba_checksum": "pass",
                },
            },
            "field_provenance": {
                "B2": {
                    "page": 1,
                    "origin": "ocr",
                    "review_required": review,
                    "review_reasons": ["ocr_origin"] if review else [],
                    "amount_cross_check": "pass",
                    "aba_checksum": "pass",
                },
            },
            "validation": validation,
            "needs_review": index == 2,
        }, filename=f"route.png [{index} of 2]")

    _run_worker(
        monkeypatch,
        tmp_path,
        _db_url,
        job.id,
        [cheque(1, "1001", True), cheque(2, "1002", False)],
    )

    response = client.get(f"/api/jobs/{job.id}/results", headers=auth["acme"])
    assert response.status_code == 200, response.text
    rows = response.json()
    assert len(rows) == 2
    assert [row["filename"] for row in rows] == [
        "Customer cheque.png [1 of 2]", "Customer cheque.png [2 of 2]"]
    assert [row["original_filename"] for row in rows] == [
        "Customer cheque.png", "Customer cheque.png"]
    assert [row["extracted_data"]["document_index"] for row in rows] == [1, 2]
    assert all(row["extracted_data"]["document_count"] == 2 for row in rows)
    assert all(row["extracted_data"]["source_pages"] == [1, 1] for row in rows)
    assert rows[0]["extracted_data"]["validation"] == common_validation
    value = rows[0]["extracted_data"]["extracted_data"]["Cheque Number"]
    provenance = rows[0]["extracted_data"]["field_provenance"]["B2"]
    for field in ("review_required", "review_reasons", "amount_cross_check",
                  "aba_checksum"):
        assert field in value
        assert field in provenance
    assert [row["needs_review"] for row in rows] == [True, True]

    exported = client.get(f"/api/jobs/{job.id}/export", headers=auth["acme"])
    assert exported.status_code == 200, exported.text
    ws = openpyxl.load_workbook(io.BytesIO(exported.content)).active
    assert ws.max_row == 3
    headers = [cell.value for cell in ws[1]]
    assert headers[:3] == ["Filename", "Page", "Cheque Number"]
    assert [ws.cell(row=n, column=1).value for n in (2, 3)] == [
        "Customer cheque.png [1 of 2]", "Customer cheque.png [2 of 2]"]
    assert [ws.cell(row=n, column=2).value for n in (2, 3)] == [1, 1]


def test_validation_warnings_use_flagged_fields_when_count_is_stale(
        monkeypatch, tmp_path, _db_url, db, users, client, auth):
    job = _job(db, users["acme"])
    _run_worker(
        monkeypatch,
        tmp_path,
        _db_url,
        job.id,
        [_result({
            "extracted_data": {},
            "validation": {
                "flagged_count": 0,
                "flagged_fields": [{
                    "field": "Amount", "reason": "amount_mismatch",
                }],
            },
        })],
    )

    response = client.get(f"/api/jobs/{job.id}/results", headers=auth["acme"])
    assert response.status_code == 200, response.text
    assert response.json()[0]["validation_warnings"]
    assert response.json()[0]["needs_review"] is True


def test_results_are_sorted_by_original_file_then_document_index(
        db, users, client, auth):
    from app.models.models import DocumentResult

    job = _job(db, users["acme"])
    for original, index in (("z.png", 2), ("a.png", 2), ("z.png", 1),
                            ("a.png", 1)):
        doc = DocumentResult(
            job_id=job.id,
            filename=f"{original} [{index} of 2]",
            original_filename=original,
            extraction_json=(
                f'{{"document_index": {index}, "document_count": 2}}'),
        )
        db.add(doc)
    db.commit()

    response = client.get(f"/api/jobs/{job.id}/results", headers=auth["acme"])
    assert response.status_code == 200, response.text
    assert [(row["original_filename"], row["extracted_data"]["document_index"])
            for row in response.json()] == [
                ("a.png", 1), ("a.png", 2), ("z.png", 1), ("z.png", 2)]
