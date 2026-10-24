"""
A demo export does not name the file it came from.

The provenance column reads `<filename> · p.N`, and the filename it uses is the
STORED name of the source document — which, on the object store production
runs, is a temp name invented by `storage.get_local_path`
(`tempfile.NamedTemporaryFile(...)`), e.g. `tmphjtbr12o.pdf`. In a demo that
column is a row of machine droppings on the thing an investor is looking at.

So for a demo tenant every export drops it, and that is FIVE separate writers,
because the filename reaches the file five different ways:

  GET  /api/jobs/{id}/export       slot writer: the Source column AND a comment
                                   on every value (both carry the filename)
  GET  /api/jobs/{id}/export       flat writer (a job with no template): the
                                   Filename column
  GET  /api/jobs/{id}/export/zip   the slot writer again, plus the NAME OF EACH
                                   FILE inside the ZIP
  POST /api/export/combined        export.py: the File column
  POST /api/export/perfile         export.py: the sheet's TAB TITLE

Every test here has a non-demo twin asserting the column is still there, since
the point is that one tenant's cosmetics must not quietly become everyone's
missing provenance.
"""
import io
import json
import zipfile

import pytest
from openpyxl import load_workbook

from tests.harness import bootstrap as bs

bs.bootstrap()

#: What the stored filename actually looks like in production, verbatim from
#: the report that prompted this.
TEMP_NAME = "tmphjtbr12o.pdf"

GRID = json.dumps({
    "cells": {"0,0": {"value": "Invoice Number", "style": {}},
              "0,1": {"value": "", "style": {}},
              "1,0": {"value": "Total", "style": {}},
              "1,1": {"value": "", "style": {}}},
    "colWidths": [160, 160], "merges": {}, "repeatRows": [],
})

SLOT_DOC = {
    "document_type": "sales_invoice",
    "template_type": "slot",
    "extraction_method": "slot_directed",
    "overall_confidence": "high",
    "extracted_fields": {"B1": "INV-1", "B2": "100.00"},
    "field_provenance": {
        "B1": {"source": "Invoice Number INV-1", "page": 1, "grounded": True},
        "B2": {"source": "Total 100.00", "page": 1, "grounded": True},
    },
    "extracted_data": {
        "Invoice Number": {"value": "INV-1", "confidence": "high", "ref": "B1"},
        "Total": {"value": "100.00", "confidence": "high", "ref": "B2"},
    },
    "slot_map": {"fields": [{"ref": "B1", "row_label": "Invoice Number"},
                            {"ref": "B2", "row_label": "Total"}],
                 "tables": []},
    "validation": {"confidence_map": {"B1": "high", "B2": "high"}},
}


def _template(db, owner, client_id, name):
    from app.models.models import ColumnTemplate
    t = db.query(ColumnTemplate).filter(ColumnTemplate.name == name).first()
    if not t:
        t = ColumnTemplate(name=name, document_type="sales_invoice",
                           description=GRID, columns_json="[]",
                           user_id=owner.id, client_id=client_id)
        db.add(t)
        db.commit()
        db.refresh(t)
    return t


def _job(db, owner, client_id, *, template=None, payload=SLOT_DOC):
    from app.models.models import DocumentResult, ExtractionJob
    j = ExtractionJob(user_id=owner.id, client_id=client_id, status="completed",
                      total_docs=1, successful=1, input_source="upload",
                      schema_id=str(template.id) if template else None)
    db.add(j)
    db.commit()
    db.refresh(j)
    db.add(DocumentResult(job_id=j.id, filename=TEMP_NAME,
                          document_type="sales_invoice",
                          overall_confidence="high", needs_review=False,
                          extraction_json=json.dumps(payload)))
    db.commit()
    return j


# ── jobs: the same content, one per tenant ───────────────────────────────────

@pytest.fixture(scope="module")
def demo_job(db, demo_user):
    tpl = _template(db, demo_user, "pearx_demo", "export_demo_tpl")
    return _job(db, demo_user, "pearx_demo", template=tpl)


@pytest.fixture(scope="module")
def plain_job(db, users):
    tpl = _template(db, users["acme"], "acme_001", "export_plain_tpl")
    return _job(db, users["acme"], "acme_001", template=tpl)


@pytest.fixture(scope="module")
def demo_job_no_template(db, demo_user):
    """A no-template job, which exports through the FLAT writer instead."""
    return _job(db, demo_user, "pearx_demo", payload={
        "extracted_data": {"Invoice Number": {"value": "INV-9",
                                              "confidence": "high"}}})


@pytest.fixture(scope="module")
def plain_job_no_template(db, users):
    return _job(db, users["acme"], "acme_001", payload={
        "extracted_data": {"Invoice Number": {"value": "INV-9",
                                              "confidence": "high"}}})


def _all_text(ws) -> str:
    """Every cell value AND every comment, as one string."""
    out = []
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None:
                out.append(str(c.value))
            if c.comment is not None:
                out.append(str(c.comment.text))
    return "\n".join(out)


def _book(content):
    return load_workbook(io.BytesIO(content))


# ── GET /api/jobs/{id}/export — the slot writer ──────────────────────────────

class TestTheTemplateExport:
    def test_a_demo_export_has_no_source_column_and_no_comments(
            self, client, demo_auth, demo_on, demo_job):
        r = client.get(f"/api/jobs/{demo_job.id}/export", headers=demo_auth)
        assert r.status_code == 200, r.text
        ws = _book(r.content).active
        text = _all_text(ws)
        assert "Source" not in text
        assert "tmphjtbr12o" not in text
        assert "· p.1" not in text
        # The extracted values are all still there — this removes a column, not
        # the export.
        assert "INV-1" in text and "100" in text

    def test_a_NON_demo_export_still_has_the_source_column(
            self, client, auth, plain_job):
        """The twin. If this ever fails, provenance was removed for everybody.

        Asserted on the LABEL rather than on a "Source" heading: the slot
        writer only writes that heading above a table band, so a fields-only
        template gets the per-row `<file> · p.N` cells with no header over
        them. Pre-existing, and not what this change is about.
        """
        r = client.get(f"/api/jobs/{plain_job.id}/export", headers=auth["acme"])
        assert r.status_code == 200, r.text
        ws = _book(r.content).active
        text = _all_text(ws)
        assert f"{TEMP_NAME} · p.1" in text, "the Source cell is gone"
        # ...and the same filename is in a comment on every value.
        assert f"{TEMP_NAME}, page 1" in text

    def test_the_demo_user_can_still_ask_for_provenance_and_not_get_it(
            self, client, demo_auth, demo_on, demo_job):
        """`?provenance=true` is the default and is overridden for a demo
        tenant: the query cannot put the temp name back."""
        r = client.get(f"/api/jobs/{demo_job.id}/export?provenance=true",
                       headers=demo_auth)
        assert "tmphjtbr12o" not in _all_text(_book(r.content).active)

    def test_with_the_variable_unset_the_same_job_exports_with_provenance(
            self, client, demo_auth, demo_job):
        """No `demo_on`: nothing about this account is a demo, so the export is
        the ordinary one."""
        r = client.get(f"/api/jobs/{demo_job.id}/export", headers=demo_auth)
        assert "tmphjtbr12o" in _all_text(_book(r.content).active)


class TestTheFlatWriter:
    def test_a_demo_export_of_a_no_template_job_has_no_filename_column(
            self, client, demo_auth, demo_on, demo_job_no_template):
        r = client.get(f"/api/jobs/{demo_job_no_template.id}/export",
                       headers=demo_auth)
        assert r.status_code == 200, r.text
        ws = _book(r.content).active
        headers = [c.value for c in ws[1]]
        assert "Filename" not in headers, headers
        # The first column is now a field, and the data moved with it.
        assert headers[0] == "Invoice Number"
        assert ws.cell(row=2, column=1).value == "INV-9"
        assert "tmphjtbr12o" not in _all_text(ws)

    def test_a_NON_demo_export_of_a_no_template_job_still_has_it(
            self, client, auth, plain_job_no_template):
        r = client.get(f"/api/jobs/{plain_job_no_template.id}/export",
                       headers=auth["acme"])
        ws = _book(r.content).active
        headers = [c.value for c in ws[1]]
        assert headers[0] == "Filename", headers
        assert ws.cell(row=2, column=1).value == TEMP_NAME


# ── GET /api/jobs/{id}/export/zip ────────────────────────────────────────────

class TestTheZipExport:
    def test_a_demo_zip_names_its_files_by_position(
            self, client, demo_auth, demo_on, demo_job):
        r = client.get(f"/api/jobs/{demo_job.id}/export/zip", headers=demo_auth)
        assert r.status_code == 200, r.text
        with zipfile.ZipFile(io.BytesIO(r.content)) as zf:
            names = zf.namelist()
            assert names == ["document_1_extracted.xlsx"], names
            inner = _book(zf.read(names[0])).active
            assert "tmphjtbr12o" not in _all_text(inner)
            assert "Source" not in _all_text(inner)

    def test_a_NON_demo_zip_still_names_its_files_after_the_document(
            self, client, auth, plain_job):
        r = client.get(f"/api/jobs/{plain_job.id}/export/zip",
                       headers=auth["acme"])
        with zipfile.ZipFile(io.BytesIO(r.content)) as zf:
            assert zf.namelist() == ["tmphjtbr12o_extracted.xlsx"]
            text = _all_text(_book(zf.read(zf.namelist()[0])).active)
            assert f"{TEMP_NAME} · p.1" in text


# ── POST /api/export/combined and /api/export/perfile ────────────────────────

class TestTheFlatExports:
    def _combined(self, client, headers, job_id):
        r = client.post("/api/export/combined", headers=headers,
                        json={"job_id": job_id})
        assert r.status_code == 200, r.text
        return _book(r.content).active

    def _perfile(self, client, headers, job_id):
        r = client.post("/api/export/perfile", headers=headers,
                        json={"job_id": job_id})
        assert r.status_code == 200, r.text
        return _book(r.content)

    def test_the_demo_combined_sheet_has_no_file_column(
            self, client, demo_auth, demo_on, demo_job):
        ws = self._combined(client, demo_auth, demo_job.id)
        headers = [c.value for c in ws[1]]
        assert "File" not in headers, headers
        assert headers[0] == "Doc Type"
        assert "tmphjtbr12o" not in _all_text(ws)
        # ...and the remaining columns did not shift out of their header.
        assert ws.cell(row=2, column=1).value == "sales_invoice"

    def test_the_NON_demo_combined_sheet_still_has_it(
            self, client, auth, plain_job):
        ws = self._combined(client, auth["acme"], plain_job.id)
        headers = [c.value for c in ws[1]]
        assert headers[0] == "File", headers
        assert ws.cell(row=2, column=1).value == TEMP_NAME

    def test_the_demo_per_file_sheet_is_not_named_after_the_temp_file(
            self, client, demo_auth, demo_on, demo_job):
        wb = self._perfile(client, demo_auth, demo_job.id)
        assert wb.sheetnames == ["Document 1", "Legend"], wb.sheetnames

    def test_the_NON_demo_per_file_sheet_is_still_named_after_the_document(
            self, client, auth, plain_job):
        wb = self._perfile(client, auth["acme"], plain_job.id)
        assert wb.sheetnames == ["tmphjtbr12o", "Legend"], wb.sheetnames
