"""Edits are addressed by SLOT, not by label.

A matrix template draws one label with two value columns: `Principal` under
`Years 1-7` and under `Years 8-30`. The label-keyed projection
(`extracted_data`) kept ONE of them, so the second value vanished from the grid
and from every label-keyed export, and an edit naming `Principal` could not say
which cell it meant. These tests use the committed hand-drawn matrix template.
"""
import json

import pytest

from tests.harness import bootstrap as bs

bs.bootstrap()

from slot_extractor import run_slot_extraction  # noqa: E402

GRID = json.loads((bs.TESTS_DIR / "gold" / "hand_drawn" /
                   "matrix_two_value_columns.json").read_text(encoding="utf-8"))
PAGE = ("Payment Calculation Years 1-7 Years 8-30\n"
        "Principal $1,000.00 $2,000.00\n"
        "Interest $50.00 $60.00\n"
        "Fees $5.00 $6.00\n"
        "Total $1,055.00 $2,066.00")
ANSWER = {"Principal": ("$1,000.00", "$2,000.00"), "Interest": ("$50.00", "$60.00"),
          "Fees": ("$5.00", "$6.00"), "Total": ("$1,055.00", "$2,066.00")}


def _extract():
    from template_shape import compute_shape
    shape = compute_shape(GRID, log=lambda _m: None)
    fields = {}
    for f in shape["field_slots"]:
        pair = ANSWER[f["row_label"]]
        v = pair[0] if f["col"] == 1 else pair[1]
        line = next(l for l in PAGE.splitlines() if l.startswith(f["row_label"]))
        fields[f["slot_id"]] = {"value": v, "source": line, "page": 1}

    class _Resp:
        success, tokens_used, model_used = True, 0, "stub"
        parsed_json = {"fields": fields}
        raw_text = json.dumps(parsed_json)

    orch = type("O", (), {"llm": type("L", (), {"extract": lambda self, **k: _Resp()})()})()
    return run_slot_extraction(orch, "M.pdf", {"layout": GRID, "shape": shape}, None,
                               page_images=[], doc_text=PAGE, doc_text_pages=[PAGE],
                               file_type="digital_pdf", default_doc_type="other",
                               start=0.0)[0].extracted_data


@pytest.fixture
def ed():
    return _extract()


class TestTwoSlotsUnderOneLabelAreBothThere:
    def test_the_projection_keeps_both_values(self, ed):
        kv = ed["extracted_data"]
        assert kv["Principal (Years 1-7)"]["value"] == "$1,000.00"
        assert kv["Principal (Years 8-30)"]["value"] == "$2,000.00"
        assert kv["Principal (Years 1-7)"]["ref"] != kv["Principal (Years 8-30)"]["ref"]
        assert len(kv) == len(ed["extracted_fields"]) == 8


class TestAnEditNamesItsSlot:
    def test_editing_one_slot_leaves_its_label_twin_alone(self, ed):
        import copy
        from app.api.routes.extract import _edit_field
        ref_a = ed["extracted_data"]["Principal (Years 1-7)"]["ref"]
        ref_b = ed["extracted_data"]["Principal (Years 8-30)"]["ref"]
        out = copy.deepcopy(ed)
        _edit_field(out, ed, ref_b, "$2,500.00")
        assert out["extracted_fields"][ref_b] == "$2,500.00"
        assert out["extracted_fields"][ref_a] == "$1,000.00"
        assert out["extracted_data"]["Principal (Years 8-30)"] == {
            "value": "$2,500.00", "confidence": "edited", "ref": ref_b}
        assert out["field_provenance"][ref_b]["original_value"] == "$2,000.00"

    def test_a_label_two_slots_share_is_refused_not_guessed(self, ed):
        import copy
        from fastapi import HTTPException
        from app.api.routes.extract import _reconcile_edits
        payload = copy.deepcopy(ed)
        payload["extracted_data"]["Principal"] = {"value": "$9.00", "confidence": "edited"}
        with pytest.raises(HTTPException) as e:
            _reconcile_edits(ed, payload)
        assert e.value.status_code == 422 and "Principal" in e.value.detail

    def test_the_grid_payload_for_a_qualified_label_finds_its_slot(self, ed):
        """The grid replaces the entry and drops its ref; the key alone must
        still resolve, through the stored projection."""
        import copy
        from app.api.routes.extract import _reconcile_edits
        ref_b = ed["extracted_data"]["Principal (Years 8-30)"]["ref"]
        payload = copy.deepcopy(ed)
        payload["extracted_data"]["Principal (Years 8-30)"] = {
            "value": "$2,500.00", "confidence": "edited"}
        out = _reconcile_edits(ed, payload)
        assert out["extracted_fields"][ref_b] == "$2,500.00"

    def test_the_export_writes_each_slot_where_it_belongs(self, ed):
        import copy
        import openpyxl
        from app.api.routes.extract import _edit_field, _write_excel
        from app.models.models import DocumentResult
        ref_b = ed["extracted_data"]["Principal (Years 8-30)"]["ref"]
        ref_a = ed["extracted_data"]["Principal (Years 1-7)"]["ref"]
        out = copy.deepcopy(ed)
        _edit_field(out, ed, ref_b, "$2,500.00")
        wb = openpyxl.Workbook()
        _write_excel(wb.active, [DocumentResult(filename="M.pdf", document_type="x",
                                                extraction_json=json.dumps(out))],
                     GRID, {}, openpyxl)
        ws = wb.active
        assert ws[ref_b].value == 2500.0 and ws[ref_a].value == 1000.0
        assert ws[ref_b].comment.text.startswith("Edited in DocAgent")
        assert "$2,000.00" in ws[ref_b].comment.text
        assert ws[ref_a].comment.text.startswith("“Principal")
