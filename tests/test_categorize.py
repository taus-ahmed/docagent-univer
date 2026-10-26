"""Categorization: rules first, the model only as a fallback (Part A2)."""
from __future__ import annotations

import io
import json
from types import SimpleNamespace

import pytest

from tests.harness import bootstrap as bs

bs.bootstrap()

from categorize import CATEGORY_UNCERTAIN, UNCATEGORIZED, categorize_rows, validate_rules  # noqa: E402

ROWS = [{"Description": "Delta Airlines ticket", "Amount": "420.00"},
        {"Description": "Uber Eats order", "Amount": "18.50"},
        {"Description": "ACME Fuel Station 4", "Amount": "60.00"},
        {"Description": "Hilton Hotel", "Amount": "199.00"}]


def _by_row(out):
    return {e["row"]: e for e in out}


def test_one_rule_decides_and_says_which():
    e = _by_row(categorize_rows(ROWS))[0]
    assert (e["category"], e["level"], e["by"], e["rule"], e["matched"]) == \
        ("Travel", "grounded", "rule", "default", "Airlines")


def test_an_ambiguous_merchant_is_not_guessed():
    e = _by_row(categorize_rows(ROWS))[1]          # "uber" (Travel) and "eats" (Meals)
    assert e["category"] == UNCATEGORIZED and CATEGORY_UNCERTAIN in e["reasons"]
    assert "Meals" in e["note"] and "Travel" in e["note"]


def test_no_rule_and_no_model_is_uncategorized_and_flagged():
    e = _by_row(categorize_rows(ROWS))[2]
    assert e["category"] == UNCATEGORIZED and e["reasons"] == [CATEGORY_UNCERTAIN]


def test_the_firm_and_template_lists_come_first():
    firm = [{"category": "Fuel", "keywords": ["fuel station"]}]
    tpl = [{"category": "Client Travel", "keywords": ["airlines"]}]
    out = _by_row(categorize_rows(ROWS, template_rules=tpl, firm_rules=firm))
    assert out[2]["category"] == "Fuel" and out[2]["rule"] == "firm"
    assert out[0]["category"] == "Client Travel" and out[0]["rule"] == "template"


def test_the_model_fills_only_what_rules_left_and_only_from_the_list():
    calls = []

    def model(texts, cats):
        calls.append(texts)
        return ["Fuel" if "Fuel" in t else "Made Up" for t in texts]
    out = _by_row(categorize_rows(ROWS, firm_rules=[{"category": "Fuel", "keywords": ["diesel"]}],
                                  model=model))
    assert calls == [["ACME Fuel Station 4"]]       # ambiguous rows are not sent
    assert (out[2]["category"], out[2]["level"], out[2]["by"]) == ("Fuel", "unverified", "model")


def test_a_model_answer_outside_the_list_is_refused():
    out = _by_row(categorize_rows(ROWS, model=lambda t, c: ["Made Up"] * len(t)))
    assert out[2]["category"] == UNCATEGORIZED and "not one of the listed" in out[2]["note"]


def test_the_model_being_down_is_said():
    def down(t, c):
        raise RuntimeError("503")
    out = _by_row(categorize_rows(ROWS, model=down))
    assert out[2]["category"] == UNCATEGORIZED and "unavailable" in out[2]["note"]


def test_rows_are_never_modified():
    rows = json.loads(json.dumps(ROWS))
    categorize_rows(rows, model=lambda t, c: ["Travel"] * len(t))
    assert rows == ROWS


def test_bad_rules_are_named_in_plain_words():
    p = validate_rules([{"category": "", "keywords": ["x"]},
                        {"category": "A"}, {"category": "B", "regex": ["(unclosed"]}])
    assert len(p) == 3 and "category name" in p[0] and "not valid" in p[2]


def _ed():
    from categorize import categorize_rows as cr
    ed = {"items_rows": [dict(r, _confidence="high") for r in ROWS],
          "slot_map": {"fields": [], "tables": [{"name": "items"}]}}
    ed["categories"] = {"items": cr(ed["items_rows"])}
    return ed


def test_categories_are_review_values_with_the_same_statuses():
    from app.core import review as R
    vals = [v for v in R.document_values(_ed()) if v["kind"] == "category"]
    by = {v["row"]: v for v in vals}
    assert by[0]["status"] == "grounded" and by[0]["resolved"]
    assert not by[1]["resolved"] and CATEGORY_UNCERTAIN in by[1]["reasons"]


def test_the_export_has_a_marked_categories_sheet():
    import openpyxl
    from app.api.routes import extract as X
    ed = _ed()
    doc = SimpleNamespace(filename="exp.pdf", id=1, get_extracted_data=lambda: ed)
    wb = openpyxl.Workbook()
    X._add_review_legend(wb, [doc])
    ws = wb["Categories"]
    values = [[c.value for c in r] for r in ws.iter_rows(min_row=2, max_row=5)]
    assert [v[3] for v in values] == ["Travel", UNCATEGORIZED, UNCATEGORIZED, "Lodging"]
    assert ws.cell(row=3, column=4).fill.fgColor.rgb != "00000000"   # flagged is filled
    assert ws.cell(row=3, column=4).comment is not None


def test_each_firm_gets_only_its_own_rules(tmp_path):
    """Isolation: the worker reads the JOB's firm schema; firm B's job never
    sees firm A's list."""
    from app.services.job_worker import _firm_categories
    a = tmp_path / "a.yaml"
    a.write_text("client_id: a\ncategories:\n  - {category: Fuel, keywords: [fuel station]}\n",
                 encoding="utf-8")
    b = tmp_path / "b.yaml"
    b.write_text("client_id: b\n", encoding="utf-8")
    assert _by_row(categorize_rows(ROWS, firm_rules=_firm_categories(a)))[2]["category"] == "Fuel"
    assert _by_row(categorize_rows(ROWS, firm_rules=_firm_categories(b)))[2]["category"] == UNCATEGORIZED
