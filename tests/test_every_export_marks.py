"""Every export entry point marks every value that is not cleanly verified.

One fixture holds a value of every status — verified, grounded, unverified,
low, edited, empty-and-flagged, unknown (a legacy "medium"), a verified value
carrying review_required, a reviewed one — each value a distinct string. Each
writer that can produce a downloaded sheet is run over it:

    _write_excel → slot writer (template_type "slot")
    _write_excel → legacy form writer (labeled, form_kv) + table-rows fallback
    _write_excel → legacy table writer (labeled, table)
    _write_excel → legacy layout writer (structural)
    _write_flat_table                  (no template, nothing inferred)
    export._build_excel per_file / combined

(`_write_inferred_sheets` draws through the slot writer; `/export/zip` calls
`_write_excel` / `_write_flat_table` per document.) For every cell holding a
fixture value whose status owes a mark, the cell must carry a fill from
EXPORT_FILLS and a comment; a clean verified value must carry neither. Every
sheet must carry a Review status column and the workbook a Legend.
"""
from __future__ import annotations

import contextlib
import io
import json
from types import SimpleNamespace

import pytest

from tests.harness import bootstrap as bs

bs.bootstrap()

import openpyxl  # noqa: E402

from app.core import review as R  # noqa: E402

# label, ref, value, level, provenance extras
FIELDS = [
    ("Verified", "B1", "val-verified", "high", {}),
    ("Grounded", "B2", "val-grounded", "grounded", {}),
    ("Unverified", "B3", "val-unverified", "unverified", {}),
    ("Low", "B4", "val-low", "low", {}),
    ("Edited", "B5", "val-edited", "edited", {"edited": True}),
    ("Unknown", "B6", "val-medium", "medium", {}),
    ("Required", "B7", "val-required", "high",
     {"review_required": True, "review_reasons": ["ocr_origin"]}),
    ("Reviewed", "B8", "val-reviewed", "low", {"reviewed": True}),
    ("Missing", "B9", "", "low", {}),
]
#: values that must be marked (anything but a clean verified/grounded one)
MARKED = {"val-unverified", "val-low", "val-edited", "val-medium",
          "val-required", "val-reviewed"}
CLEAN = {"val-verified", "val-grounded"}

ROWS = [
    {"Item": "row-high", "Amount": "row-amt-high", "_confidence": "high"},
    {"Item": "row-low", "Amount": "row-amt-low", "_confidence": "high",
     "_ungrounded": ["Item", "Amount"]},
    {"Item": "row-legacy", "Amount": "row-amt-legacy"},       # no level: unknown
]
ROW_MARKED = {"row-low", "row-amt-low", "row-legacy", "row-amt-legacy"}
ROW_CLEAN = {"row-high", "row-amt-high"}


def _result(template_type="slot"):
    ed = {
        "template_type": template_type,
        "extracted_fields": {ref: v for _, ref, v, _, _ in FIELDS},
        "extracted_data": {lab: {"value": v, "confidence": lvl, "ref": ref}
                           for lab, ref, v, lvl, _ in FIELDS},
        "field_provenance": {ref: {"source": v, "page": 1, "grounded": True, **x}
                             for _, ref, v, _, x in FIELDS},
        "validation": {"confidence_map": {ref: lvl for _, ref, _, lvl, _ in FIELDS},
                       "flagged_fields": [{"ref": "B9", "value": "",
                                           "reason": "expected a value"}]},
        "slot_map": {"fields": [{"ref": ref, "row_label": lab, "row": i, "col": 1}
                                for i, (lab, ref, *_r) in enumerate(FIELDS)],
                     "tables": []},
        "table_rows": [dict(r) for r in ROWS],
    }
    return ed


def _doc(ed, name="D.pdf"):
    s = json.dumps(ed)
    return SimpleNamespace(filename=name, document_type="invoice", id=1,
                           overall_confidence="low", needs_review=True,
                           get_extracted_data=lambda: json.loads(s))


def _grid(extract_target=False):
    cells = {}
    for i, (lab, *_r) in enumerate(FIELDS):
        cells[f"{i},0"] = {"value": lab}
        cells[f"{i},1"] = {"value": "", **({"extractTarget": True} if extract_target else {})}
    return {"cells": cells, "colWidths": [], "merges": {}}


def _cells(wb):
    for ws in wb.worksheets:
        if ws.title == "Legend":
            continue
        for row in ws.iter_rows():
            for c in row:
                if c.value not in (None, ""):
                    yield ws, c


def _check(wb, marked, clean):
    fills = set(R.EXPORT_FILLS.values())
    seen = set()
    for ws, c in _cells(wb):
        v = str(c.value)
        if v in marked:
            seen.add(v)
            assert c.fill.fgColor.rgb in fills and c.comment is not None, \
                (ws.title, c.coordinate, v, c.fill.fgColor.rgb)
        elif v in clean:
            seen.add(v)
            assert c.fill.fgColor.rgb not in fills, (ws.title, c.coordinate, v)
    assert "Legend" in wb.sheetnames
    for ws in wb.worksheets:
        if ws.title in ("Legend", "Summary"):
            continue
        heads = [str(c.value) for row in ws.iter_rows() for c in row
                 if str(c.value or "").startswith("Review status")]
        assert heads, f"{ws.title}: no Review status column"
    return seen


def _excel(ed, grid, regions=None):
    from app.api.routes import extract as X
    wb = openpyxl.Workbook()
    docs = [_doc(ed)]
    with contextlib.redirect_stdout(io.StringIO()):
        X._write_excel(wb.active, docs, grid, regions, openpyxl)
        X._add_review_legend(wb, docs)
    return wb


def test_the_status_lookup_matches_the_screen():
    """field_status / row_cell_status agree with document_values."""
    ed = _result()
    by_ref = {v["ref"]: v for v in R.document_values(ed) if v["kind"] == "field"}
    for lab, ref, *_r in FIELDS:
        a, b = R.field_status(ed, ref=ref), by_ref[ref]
        assert (a["status"], a["resolved"]) == (b["status"], b["resolved"]), ref
    cells = [v for v in R.document_values(ed) if v["kind"] == "cell"]
    for v in cells:
        st = R.row_cell_status(ed, ed["table_rows"][v["row"]], v["column"])
        assert (st["status"], st["resolved"]) == (v["status"], v["resolved"])


def test_an_image_upload_is_never_verified():
    ed = _result("labeled")
    ed["image_upload"] = True
    assert R.field_status(ed, ref="B1")["status"] == "unverified"
    assert R.field_status(ed, ref="B5")["status"] == "edited"
    assert R.row_cell_status(ed, ed["table_rows"][0], "Item")["status"] == "unverified"


def test_slot_writer():
    seen = _check(_excel(_result("slot"), _grid()), MARKED, CLEAN)
    assert seen >= MARKED | CLEAN


def test_legacy_form_writer_and_its_table_rows():
    ed = _result("labeled")
    seen = _check(_excel(ed, _grid(True), {"primary_mode": "form_kv"}),
                  MARKED | ROW_MARKED, CLEAN | ROW_CLEAN)
    assert seen >= MARKED | CLEAN | ROW_MARKED | ROW_CLEAN


def test_legacy_table_writer():
    ed = _result("labeled")
    grid = {"cells": {"0,0": {"value": "Item"}, "0,1": {"value": "Amount"}},
            "colWidths": [], "merges": {}}
    regions = {"primary_mode": "table", "table_regions": [
        {"header_row": 0, "start_col": 0, "column_names": ["Item", "Amount"]}]}
    seen = _check(_excel(ed, grid, regions), ROW_MARKED, ROW_CLEAN)
    assert seen >= ROW_MARKED | ROW_CLEAN


def test_legacy_layout_writer():
    ed = _result("structural")
    ed["extracted_fields"] = {}
    ed["layout_sections"] = {"Assets": {"rows": [
        {"label": "Cash", "value": "lay-low", "confidence": "low",
         "row": 2, "label_col": "A", "value_col": "B"},
        {"label": "Bank", "value": "lay-high", "confidence": "high",
         "row": 3, "label_col": "A", "value_col": "B"},
        {"label": "Misc", "value": "lay-none",
         "row": 4, "label_col": "A", "value_col": "B"}]}}
    grid = {"cells": {"0,0": {"value": "Assets"}}, "colWidths": [], "merges": {}}
    wb = _excel(ed, grid)
    seen = _check(wb, {"lay-low", "lay-none"}, {"lay-high"})
    assert seen == {"lay-low", "lay-none", "lay-high"}
    legend = wb["Legend"]["A1"].value
    assert legend.startswith(str(R.summarise(R.export_values(ed))["needs_review"]))


def test_flat_table():
    from app.api.routes import extract as X
    wb = openpyxl.Workbook()
    docs = [_doc(_result("slot"))]
    X._write_flat_table(wb.active, docs, openpyxl)
    X._add_review_legend(wb, docs)
    seen = _check(wb, MARKED, CLEAN)
    assert seen >= MARKED | CLEAN


@pytest.mark.parametrize("per_file", [True, False])
def test_combined_and_per_file(per_file):
    from app.api.routes.export import _build_excel
    job = SimpleNamespace(id=1, total_docs=1, successful=1, failed=0,
                          needs_review=1, total_time_sec=1.0)
    data = _build_excel(job, [_doc(_result("slot"))], None, per_file=per_file)
    wb = openpyxl.load_workbook(io.BytesIO(data))
    seen = _check(wb, MARKED, CLEAN)
    assert seen >= MARKED | CLEAN


# ── Run 15 item 8: every writer for real, exact counts, and a mutation proof ──
#
# Each writer below is run over the fixture and must (1) mark every value
# owing review and leave clean ones clean (`_check`) and (2) head its Review
# status column with EXACTLY the number of unresolved values it wrote. The
# fixture's unresolved values are 4 fields (unverified, low, medium,
# required), B9 (empty and flagged) where the writer writes it, and 4 table
# cells (row-low and row-legacy, both columns).
#
# MUTATION: the same check runs again with that ONE writer's marks disabled
# (every other writer untouched), and must fail. A test that still passes
# with the marks off proves nothing about them.

def _mixed_grid():
    grid = _grid(True)
    grid["cells"]["10,0"] = {"value": "Item"}
    grid["cells"]["10,1"] = {"value": "Amount"}
    return grid


MIXED_REGIONS = {"primary_mode": "mixed", "table_regions": [
    {"header_row": 10, "start_col": 0, "column_names": ["Item", "Amount"],
     "section_label": ""}]}
TABLE_GRID = {"cells": {"0,0": {"value": "Item"}, "0,1": {"value": "Amount"}},
              "colWidths": [], "merges": {}}
TABLE_REGIONS = {"primary_mode": "table", "table_regions": [
    {"header_row": 0, "start_col": 0, "column_names": ["Item", "Amount"]}]}


def _layout_result():
    ed = _result("structural")
    ed["extracted_fields"] = {}
    ed["layout_sections"] = {"Assets": {"rows": [
        {"label": "Cash", "value": "lay-low", "confidence": "low",
         "row": 2, "label_col": "A", "value_col": "B"},
        {"label": "Bank", "value": "lay-high", "confidence": "high",
         "row": 3, "label_col": "A", "value_col": "B"},
        {"label": "Misc", "value": "lay-none",
         "row": 4, "label_col": "A", "value_col": "B"}]}}
    return ed


def _off(monkeypatch, name):
    """Disable the marks of the writer `name` only."""
    import app.services.excel_slot as XS
    from app.core.export_marks import SheetMarks
    if name == "slot":
        real = XS._write_slot_excel
        monkeypatch.setattr(XS, "_write_slot_excel",
                            lambda *a, **k: real(*a, **{**k, "marks": False}))
        return
    attr = {"form": "_write_form_excel", "table": "_write_table_excel",
            "mixed": "_write_mixed_excel", "layout": "_write_layout_excel"}[name]
    real = getattr(XS, attr)

    def unmarked(ws, *a, **k):
        k["marks"] = SheetMarks(ws, enabled=False)
        return real(ws, *a, **k)
    monkeypatch.setattr(XS, attr, unmarked)


def _run_writer(name, monkeypatch, disabled=False):
    from app.api.routes import extract as X
    if name in ("combined", "per_file"):
        from app.api.routes import export as E
        if disabled:
            from app.core.export_marks import SheetMarks
            monkeypatch.setattr(E, "SheetMarks",
                                lambda ws: SheetMarks(ws, enabled=False),
                                raising=False)
            import app.core.export_marks as EM
            monkeypatch.setattr(EM, "SheetMarks",
                                lambda ws, enabled=True: SheetMarks(ws, enabled=False))
        job = SimpleNamespace(id=1, total_docs=1, successful=1, failed=0,
                              needs_review=1, total_time_sec=1.0)
        data = E._build_excel(job, [_doc(_result("slot"))], None,
                              per_file=(name == "per_file"))
        return openpyxl.load_workbook(io.BytesIO(data))
    if name == "flat":
        wb = openpyxl.Workbook()
        docs = [_doc(_result("slot"))]
        with contextlib.redirect_stdout(io.StringIO()):
            X._write_flat_table(wb.active, docs, openpyxl, marks=not disabled)
            X._add_review_legend(wb, docs)
        return wb
    if disabled:
        _off(monkeypatch, name)
    if name == "slot":
        return _excel(_result("slot"), _grid())
    if name == "form":
        return _excel(_result("labeled"), _grid(True), {"primary_mode": "form_kv"})
    if name == "table":
        return _excel(_result("labeled"), TABLE_GRID, TABLE_REGIONS)
    if name == "mixed":
        return _excel(_result("labeled"), _mixed_grid(), MIXED_REGIONS)
    if name == "layout":
        return _excel(_layout_result(), {"cells": {"0,0": {"value": "Assets"}},
                                         "colWidths": [], "merges": {}})
    raise KeyError(name)


#: writer -> (values that must be marked, values that must be clean,
#:            the exact unresolved count its Review status header states)
WRITERS = {
    "slot": (MARKED, CLEAN, 5),
    "form": (MARKED | ROW_MARKED, CLEAN | ROW_CLEAN, 8),
    "table": (ROW_MARKED, ROW_CLEAN, 4),
    "mixed": (MARKED | ROW_MARKED, CLEAN | ROW_CLEAN, 8),
    # 3, not 2: the legacy layout writer writes this fixture's section twice
    # (B2-B4, then B3-B5), so lay-low sits in TWO cells and the header counts
    # cells. Legacy jobs only; recorded in PROGRESS (Run 15), not changed here.
    "layout": ({"lay-low", "lay-none"}, {"lay-high"}, 3),
    "flat": (MARKED, CLEAN, 5),
    "combined": (MARKED, CLEAN, 5),
    "per_file": (MARKED, CLEAN, 5),
}


def _stated_counts(wb):
    import re
    out = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = str(c.value or "")
                if v.startswith("Review status"):
                    m = re.search(r"(\d+) values? needs? review", v)
                    out.append(int(m.group(1)) if m else 0)
    return out


def _full_check(wb, name):
    marked, clean, count = WRITERS[name]
    seen = _check(wb, marked, clean)
    assert seen >= marked | clean, (name, (marked | clean) - seen)
    assert _stated_counts(wb) == [count], (name, _stated_counts(wb))


@pytest.mark.parametrize("name", sorted(WRITERS))
def test_every_writer_marks_and_states_the_exact_count(name, monkeypatch):
    _full_check(_run_writer(name, monkeypatch), name)


@pytest.mark.parametrize("name", sorted(WRITERS))
def test_mutation_each_writers_marks_off_fails_the_check(name, monkeypatch):
    wb = _run_writer(name, monkeypatch, disabled=True)
    with pytest.raises(AssertionError):
        _full_check(wb, name)


def test_the_mixed_writer_is_the_one_that_ran(monkeypatch):
    """The mixed case must not quietly route to another writer."""
    import app.services.excel_slot as XS
    calls = []
    real = XS._write_mixed_excel
    monkeypatch.setattr(XS, "_write_mixed_excel",
                        lambda *a, **k: calls.append(1) or real(*a, **k))
    _run_writer("mixed", monkeypatch)
    assert calls == [1]
