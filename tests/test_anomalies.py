"""Rule-based anomalies with evidence (Part A3); clean data gives none."""
from __future__ import annotations

import pytest

from tests.harness import bootstrap as bs

bs.bootstrap()

from anomalies import CAP, detect  # noqa: E402


def _rows(spec):
    return [{"Date": d, "Description": t, "Amount": a} for d, t, a in spec]


CLEAN = _rows([("2024-03-04", "Office rent March", "1,250.00"),
               ("2024-03-05", "Electricity", "182.40"),
               ("2024-03-06", "Printer toner", "64.99"),
               ("2024-03-07", "Coffee beans", "23.10"),
               ("2024-03-11", "Software licence", "49.00"),
               ("2024-03-12", "Courier", "17.85"),
               ("2024-03-13", "Stationery", "31.20"),
               ("2024-03-14", "Internet", "59.99"),
               ("2024-03-15", "Train tickets", "88.40")])


def _ed(rows, **extra):
    return {"items_rows": rows, **extra}


def _rules(ed):
    return [a["rule"] for a in detect(ed)]


def test_clean_data_gives_nothing():
    assert detect(_ed(CLEAN)) == []


def test_exact_and_near_duplicates():
    rows = CLEAN + _rows([("2024-03-05", "Electricity", "182.40"),
                          ("2024-03-06", "Toner refill", "64.99"),
                          ("2024-03-14", "Internet", "59.99")])
    rows[-1]["Date"] = "2024-03-16"
    got = detect(_ed(rows))
    dup = [a for a in got if a["rule"] == "duplicate_row"]
    assert dup and dup[0]["rows"] == ["items_rows[1]", "items_rows[9]"]
    assert dup[0]["numbers"]["amount"] == "182.40"
    near = [a for a in got if a["rule"] == "near_duplicate_row"]
    assert len(near) == 2


def _courier(n, big=None):
    rows = _rows([(f"2024-03-{4 + d:02d}", "Courier", f"{17 + d}.50") for d in range(n)])
    if big:
        rows.append({"Date": "2024-03-28", "Description": "Courier", "Amount": big})
    cats = {"items": [{"row": i, "category": "Postage", "level": "grounded"} for i in range(len(rows))]}
    return _ed(rows, categories=cats)


def test_an_outlier_within_a_category_needs_enough_values():
    out = [a for a in detect(_courier(9, "9,850.00")) if a["rule"] == "outlier_amount"]
    assert out and out[0]["rows"] == ["items_rows[9]"] and "median" in out[0]["numbers"]
    assert "Postage" in out[0]["message"]
    assert not [a for a in detect(_courier(4, "9,850.00")) if a["rule"] == "outlier_amount"]


def test_mixed_kinds_of_spending_are_not_outliers():
    """Rent beside stationery: uncategorized rows are never compared."""
    rows = CLEAN + _rows([("2024-03-18", "Courier", "9,850.00")])
    assert "outlier_amount" not in _rules(_ed(rows))


def test_round_amount_cluster():
    rows = _rows([(f"2024-03-{d:02d}", f"Item {d}", f"{100 * d}.00") for d in range(4, 14)])
    assert "round_amounts" in _rules(_ed(rows))


def test_weekend_and_out_of_period_dates():
    rows = CLEAN + _rows([("2024-03-09", "Hardware store", "41.00")])   # a Saturday
    assert "weekend_date" in _rules(_ed(rows))
    fields = {"slot_map": {"fields": [{"ref": "B1", "row_label": "Period From"},
                                      {"ref": "B2", "row_label": "Period To"}]},
              "extracted_fields": {"B1": "2024-03-01", "B2": "2024-03-15"}}
    rows = CLEAN + _rows([("2024-04-02", "Late entry", "12.00")])
    out = [a for a in detect(_ed(rows, **fields)) if a["rule"] == "out_of_period_date"]
    assert out and out[0]["numbers"]["to"] == "2024-03-15"


def test_sequence_gaps_and_repeats():
    rows = [{"Cheque No": str(n), "Amount": "10.00", "Description": f"c{n}"}
            for n in (1001, 1002, 1002, 1006, 1007)]
    rules = _rules(_ed(rows))
    assert "sequence_repeat" in rules and "sequence_gap" in rules


def test_a_zero_in_a_positive_column():
    rows = CLEAN + _rows([("2024-03-19", "Refund adjustment", "0.00")])
    assert "unusual_non_positive" in _rules(_ed(rows))


def test_failed_totals_and_currency_mix_are_reported():
    ed = _ed(CLEAN, field_provenance={"B9": {"review_reasons": ["totals_mismatch"]}},
             extracted_fields={"B1": "£10.00", "B2": "€5.00"})
    rules = _rules(ed)
    assert "totals_mismatch" in rules and "currency_mix" in rules


def test_capped_and_sorted_by_severity():
    rows = [{"Cheque No": str(1000 + 3 * n), "Amount": "10.00", "Description": "same"}
            for n in range(120)]
    got = detect(_ed(rows))
    assert len(got) == CAP + 1 and got[-1]["rule"] == "more"
    sev = [a["severity"] for a in got[:-1]]
    assert sev == sorted(sev, key=["high", "medium", "low"].index)


@pytest.mark.parametrize("name", ["INV-2024-0031", "STMT-2024-01", "EXP-2024-0081",
                                  "PO-2024-0018", "IS-2024-Q4"])
def test_the_gold_results_give_no_false_alarms(name):
    """The stored gold extraction results are clean documents."""
    import json
    rep = json.loads((bs.TESTS_DIR / "reports" / "latest.json").read_text(encoding="utf-8"))
    doc = rep["documents"][name]
    adapted = doc.get("adapted") or {}
    ed = {k: v for k, v in adapted.items() if isinstance(v, list)}
    assert [a for a in detect(ed) if a["severity"] == "high"] == []


def test_exported_on_their_own_sheet():
    import openpyxl
    from types import SimpleNamespace
    from app.api.routes import extract as X
    rows = CLEAN + _rows([("2024-03-05", "Electricity", "182.40")])
    ed = _ed(rows)
    ed["anomalies"] = detect(ed)
    wb = openpyxl.Workbook()
    X._add_review_legend(wb, [SimpleNamespace(filename="s.pdf", id=1, get_extracted_data=lambda: ed)])
    ws = wb["Anomalies"]
    assert ws.cell(row=2, column=3).value == "duplicate_row"
    assert "items_rows[1]" in ws.cell(row=2, column=5).value
