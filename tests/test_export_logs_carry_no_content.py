"""The export writers log ids and counts, never template text or values
(Run 15, item 7).

The legacy layout writer logged the first 30 characters of a static template
cell it refused to overwrite — a firm's own label text in the server log.
Every writer is run here over a template and a result whose every string
carries a sentinel, with stdout and every logger captured.
"""
from __future__ import annotations

import contextlib
import io
import logging

import openpyxl
import pytest

from tests.harness import bootstrap as bs

bs.bootstrap()

from tests.test_every_export_marks import _doc, _result  # noqa: E402

S = "QZXTPL"


def _sentinel_grid(extract_target=False):
    cells = {f"{i},0": {"value": f"{S}label{i}"} for i in range(9)}
    for i in range(9):
        cells[f"{i},1"] = {"value": "", **({"extractTarget": True} if extract_target else {})}
    cells["0,2"] = {"value": f"{S}header"}
    return {"cells": cells, "colWidths": [], "merges": {}}


def _sentinel_result(kind):
    ed = _result(kind)
    ed["extracted_fields"] = {k: f"{S}val{k}" for k in ed["extracted_fields"]}
    ed["extracted_fields"]["A1"] = f"{S}onto-a-label"   # targets a static cell
    ed["extracted_fields"]["C1"] = f"{S}onto-a-header"
    ed["table_rows"] = [{f"{S}Item": f"{S}row", "Amount": f"{S}amt"}]
    ed["layout_sections"] = {f"{S}Assets": {"rows": [
        {"label": f"{S}Cash", "value": f"{S}lay", "confidence": "low",
         "row": 2, "label_col": "A", "value_col": "B"}]}}
    return ed


CASES = {
    "slot": ("slot", False, None),
    "structural": ("structural", False, None),
    "form": ("labeled", True, {"primary_mode": "form_kv"}),
    "mixed": ("labeled", True, {"primary_mode": "mixed", "table_regions": [
        {"header_row": 0, "start_col": 2, "column_names": [f"{S}header"]}]}),
    "table": ("labeled", False, {"primary_mode": "table", "table_regions": [
        {"header_row": 0, "start_col": 0, "column_names": [f"{S}label0"]}]}),
}


@pytest.mark.parametrize("case", sorted(CASES))
def test_no_writer_logs_template_text_or_values(case, capfd, caplog):
    from app.api.routes import extract as X
    kind, target, regions = CASES[case]
    caplog.set_level(logging.DEBUG)
    out = io.StringIO()
    wb = openpyxl.Workbook()
    with contextlib.redirect_stdout(out):
        X._write_excel(wb.active, [_doc(_sentinel_result(kind))],
                       _sentinel_grid(target), regions, openpyxl)
    logged = out.getvalue() + "".join(capfd.readouterr()) + caplog.text
    assert S not in logged, [ln for ln in logged.splitlines() if S in ln][:5]


def test_the_layout_writer_still_reports_the_skip_by_ref(capfd):
    """Ids and counts stay: the skipped cell is named by its reference."""
    from app.api.routes import extract as X
    out = io.StringIO()
    wb = openpyxl.Workbook()
    with contextlib.redirect_stdout(out):
        X._write_excel(wb.active, [_doc(_sentinel_result("structural"))],
                       _sentinel_grid(), None, openpyxl)
    assert "SKIP extracted_fields A1 -> static template cell" in out.getvalue()


def test_the_vision_prompt_logs_counts_not_template_text(capsys):
    """The image fallback's prompt builder logged each table's section label,
    its column names and 300 characters of the prompt itself."""
    from app.services.vision_prompts import _build_vision_prompt
    template_data = {"doc_type": "invoice", "layout": {}, "regions": {
        "primary_mode": "table",
        "table_regions": [{"section_label": f"{S}Lines", "header_row": 0,
                           "start_col": 0,
                           "column_names": [f"{S}Item", f"{S}Amount"]}]}}
    _build_vision_prompt(template_data, doc_text="")
    out = capsys.readouterr().out
    assert "[PROMPT] table_1: 2 column(s)" in out
    assert S not in out, [ln for ln in out.splitlines() if S in ln]


def test_the_layout_writer_logs_no_section_heading():
    """Section headings are the document's (or the template's) own text: the
    dedup and the semantic-match lines say what happened, not to what."""
    from app.api.routes import extract as X
    rows = [{"label": "Loan", "value": "1", "confidence": "high",
             "row": 2, "label_col": "A", "value_col": "B"}]
    ed = _result("structural")
    ed["extracted_fields"] = {}
    ed["layout_sections"] = {"Long-term Liabilities": {"rows": rows},
                             "Long-term Liabilities again": {"rows": list(rows)}}
    ed["binding_column_groups"] = [{"section_label": "Non current liabilities", "group_id": 0,
                                    "label_col_letter": "A",
                                    "value_col_letter": "B", "start_row": 1}]
    out = io.StringIO()
    wb = openpyxl.Workbook()
    with contextlib.redirect_stdout(out):
        X._write_excel(wb.active, [_doc(ed)],
                       {"cells": {"0,0": {"value": "Non current liabilities"}},
                        "colWidths": [], "merges": {}}, None, openpyxl)
    log = out.getvalue()
    assert "dropped a section" in log and "semantic section match" in log, log
    assert "liabilit" not in log.lower(), log


def test_row_alignment_logs_no_table_or_column_name(capsys):
    from app.services.result_processing import _validate_row_alignment
    regions = [{"section_label": f"{S}Lines",
                "column_names": [f"{S}Item", f"{S}Amount"]}]
    rows = ([{f"{S}Item": "", f"{S}Amount": "", "_table_source": "table"}] * 3
            + [{f"{S}Item": "x", f"{S}Amount": "1", "_table_source": "table"}] * 2
            + [{f"{S}Item": "x", "_table_source": "table"}] * 3)
    _rows, misaligned, warning = _validate_row_alignment(rows, regions)
    out = capsys.readouterr().out
    assert "[BUG6]" in out
    assert S not in out, out
