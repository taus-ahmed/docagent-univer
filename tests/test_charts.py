"""Charts from exact aggregations, one currency at a time (Part A5)."""
from __future__ import annotations

import json
from types import SimpleNamespace

from tests.harness import bootstrap as bs

bs.bootstrap()

from charts import aggregate  # noqa: E402


def _ed(rows, **kw):
    return {"items_rows": rows, **kw}


ROWS = [{"Date": "2024-03-04", "Description": "Rent", "Amount": "$1,250.10"},
        {"Date": "2024-03-19", "Description": "Courier", "Amount": "$0.20"},
        {"Date": "2024-04-02", "Description": "Rent", "Amount": "$1,250.10"},
        {"Date": "2024-04-03", "Description": "Hotel", "Amount": "€99.99"}]


def test_sums_are_exact_decimals():
    g = aggregate([_ed(ROWS)])["currencies"]
    usd = next(v for k, v in g.items() if "USD" in k)
    assert dict(usd["by_month"]) == {"2024-03": "1250.30", "2024-04": "1250.10"}
    assert usd["top_items"][0] == ["Rent", "2500.20"]


def test_currencies_are_never_added_together():
    out = aggregate([_ed(ROWS)])
    assert set(out["currencies"]) == {"USD/CAD/AUD/SGD/HKD/NZD/MXN", "EUR"}
    assert out["currencies"]["EUR"]["by_month"] == [["2024-04", "99.99"]]
    assert "2 currencies" in out["message"]


def test_ambiguous_amounts_are_left_out_and_counted():
    # no number on the record says which style it uses: "1,234" is not charted
    out = aggregate([_ed([{"Description": "x", "Amount": "1,234"},
                          {"Description": "y", "Amount": "7"},
                          {"Description": "z", "Amount": "2"}])])
    assert out["left_out_ambiguous"] == 1 and out["values"] == 2
    # ...while a document whose other numbers are "7.50" resolves it
    assert aggregate([_ed([{"Description": "x", "Amount": "1,234"},
                           {"Description": "y", "Amount": "7.50"}])])["left_out_ambiguous"] == 0


def test_too_little_data_says_so():
    out = aggregate([_ed([{"Description": "x", "Amount": "7.50"}])])
    assert "Not enough" in out["message"]
    assert aggregate([])["message"]


def test_categories_drive_spend_by_category_and_edits_are_used():
    ed = _ed([{"Description": "Rent", "Amount": "100.00"},
              {"Description": "Taxi", "Amount": "20.00", "_edited": ["Amount"]}],
             categories={"items": [{"row": 0, "category": "Rent"}, {"row": 1, "category": "Travel"}]})
    g = next(iter(aggregate([ed])["currencies"].values()))
    assert dict(g["by_category"]) == {"Rent": "100.00", "Travel": "20.00"}


def test_the_workbook_gets_native_charts_only_when_graphs_were_asked():
    import openpyxl
    from app.api.routes import extract as X
    ed = _ed(ROWS, categories={"items": [{"row": i, "category": "Ops"} for i in range(4)]},
             options_requested=["graphs"])
    doc = SimpleNamespace(filename="a.pdf", id=1, get_extracted_data=lambda: json.loads(json.dumps(ed)))
    wb = openpyxl.Workbook()
    X._add_review_legend(wb, [doc])
    ws = wb["Charts"]
    assert len(ws._charts) == 4                      # 2 currencies x (category, month)
    assert "need review" in ws["A1"].value or "checked" in ws["A1"].value
    titles = [c.title.tx.rich.p[0].r[0].t for c in ws._charts]
    assert any("(EUR)" in t for t in titles) and not any("+" in t for t in titles)
    del ed["options_requested"]
    wb2 = openpyxl.Workbook()
    X._add_review_legend(wb2, [SimpleNamespace(filename="a.pdf", id=1, get_extracted_data=lambda: ed)])
    assert "Charts" not in wb2.sheetnames
