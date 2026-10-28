"""The slot writer, the inferred-sheet writer and the export dispatcher `_write_excel`.

Moved verbatim from app/api/routes/extract.py (round 13); extract.py re-exports every name.
"""

import app.services.paths  # noqa: F401  (engine on sys.path)
from app.services.excel_cells import _apply_cell_style
from app.services.excel_cells import _apply_template_merges
from app.services.excel_cells import _find_template_dimensions
from app.services.excel_cells import _fit_columns
from app.services.excel_cells import _ref_to_rowcol
from app.services.excel_legacy import _write_form_excel
from app.services.excel_legacy import _write_layout_excel
from app.services.excel_legacy import _write_mixed_excel
from app.services.excel_legacy import _write_table_excel
from app.services.excel_cells import cell_format
from app.services.excel_cells import coerce_cell_value
from app.services.excel_cells import keep_precision_note
from app.services.excel_cells import precision_flag
from app.services.excel_cells import provenance_edited_note
from app.services.excel_cells import provenance_note
from app.services.excel_cells import provenance_note_size
import re
# --imports-end--


def _write_excel(ws, doc_results, sheet_data, template_regions, openpyxl_mod,
                 provenance=True, marks=True):
    """Route to the correct writer based on template regions primary_mode.
    Always uses the template structure for routing — never relies on AI output
    flags which can be stale or wrong from a previous extraction run."""
    cells_tpl = sheet_data.get("cells", {})
    col_widths = sheet_data.get("colWidths", [])
    max_r, max_c = _find_template_dimensions(cells_tpl)

    # Deterministic writer routing that mirrors the extraction routing.
    # The authoritative signal is the template_type saved in extraction_json:
    #   structural        -> layout writer
    #   labeled | mixed   -> form/mixed/table writer (by primary_mode)
    #   (no template/flat is handled by the export endpoint via _write_flat_table)
    # Legacy jobs (extracted before template_type was persisted) carry no
    # template_type — fall back to the previous behaviour (layout_sections present
    # -> layout writer) so already-extracted balance sheets still export correctly.
    def _has_layout(d):
        ed = d.get_extracted_data()
        ls = ed.get("layout_sections") if isinstance(ed, dict) else None
        return isinstance(ls, dict) and any(
            isinstance(s, dict) and s.get("rows") for s in ls.values()
        )

    template_type = ""
    for d in doc_results:
        ed = d.get_extracted_data()
        if isinstance(ed, dict) and ed.get("template_type"):
            template_type = ed["template_type"]
            break

    primary_mode = (template_regions or {}).get("primary_mode", "form_kv")

    if template_type == "slot":
        print("[EXPORT] routing: template_type=slot -> slot writer "
              "(values written to the addresses they were requested for)", flush=True)
        _write_slot_excel(ws, doc_results, sheet_data, cells_tpl, openpyxl_mod,
                          provenance=provenance, marks=marks)
        _fit_columns(ws, col_widths)
        return
    # Every other writer marks through one SheetMarks (app/core/export_marks):
    # the same fills, comments and Review status column as the slot writer.
    from app.core.export_marks import SheetMarks
    sm = SheetMarks(ws, enabled=marks)
    if template_type == "structural":
        print("[EXPORT] routing: template_type=structural -> layout-mode writer", flush=True)
        _write_layout_excel(ws, doc_results, sheet_data, cells_tpl, openpyxl_mod,
                            marks=sm)
        sm.finish()
        _fit_columns(ws, col_widths)
        return
    if template_type in ("labeled", "mixed"):
        # field path — never the layout writer, even if a stale layout_sections exists
        print(f"[EXPORT] routing: template_type={template_type} -> "
              f"form/mixed/table writer (primary_mode={primary_mode})", flush=True)
    else:
        # legacy (no template_type) — preserve prior layout-detection behaviour
        if any(_has_layout(d) for d in doc_results):
            print("[EXPORT] routing: layout-mode writer (legacy layout_sections)", flush=True)
            _write_layout_excel(ws, doc_results, sheet_data, cells_tpl, openpyxl_mod,
                                marks=sm)
            sm.finish()
            _fit_columns(ws, col_widths)
            return
        print(f"[EXPORT] routing: legacy primary_mode={primary_mode}", flush=True)

    if primary_mode == "table":
        _write_table_excel(ws, doc_results, sheet_data, cells_tpl, template_regions, openpyxl_mod,
                           marks=sm)
    elif primary_mode == "mixed":
        _write_mixed_excel(ws, doc_results, sheet_data, cells_tpl, max_r, max_c, template_regions, openpyxl_mod,
                           marks=sm)
    else:
        # form_with_targets, form_kv, two_column — all use form writer
        _write_form_excel(ws, doc_results, sheet_data, cells_tpl, max_r, max_c, openpyxl_mod,
                          marks=sm)
    sm.finish()
    _fit_columns(ws, col_widths)


def _write_inferred_sheets(wb, ws, doc_results, openpyxl_mod, provenance=True):
    """Phase 3 — export a job that ran with no template, using the template the
    engine inferred for each document.

    Documents whose inferred shape matches share ONE sheet and stack in it;
    genuinely different shapes get their own sheet. Never one sheet per
    document when the shapes match. Returns False if nothing was inferred.
    """
    groups = {}
    for d in doc_results:
        ed = d.get_extracted_data()
        if not isinstance(ed, dict):
            continue
        grid = ed.get("inferred_grid")
        if not (isinstance(grid, dict) and grid.get("cells")):
            continue
        sig = ed.get("shape_signature") or "shape"
        groups.setdefault(sig, {"grid": grid, "docs": [],
                                "title": (ed.get("inferred_template") or {}).get("title", "")})
        groups[sig]["docs"].append(d)
    if not groups:
        return False

    used_titles = set()

    def _sheet_title(raw, idx):
        base = re.sub(r"[\\/*?:\[\]]", " ", str(raw or "")).strip() or f"Sheet {idx}"
        base = base[:28] or f"Sheet {idx}"
        title, n = base, 2
        while title.casefold() in used_titles:
            title = f"{base[:25]} {n}"
            n += 1
        used_titles.add(title.casefold())
        return title

    print(f"[EXPORT] no template: {len(doc_results)} document(s) share "
          f"{len(groups)} inferred shape(s)", flush=True)

    for i, (sig, g) in enumerate(groups.items()):
        sheet = ws if i == 0 else wb.create_sheet()
        sheet.title = _sheet_title(g["title"], i + 1)
        # the sheet's title is a template or document title: its index only
        print(f"[EXPORT]   shape {sig}: {len(g['docs'])} document(s) "
              f"-> sheet {i + 1}", flush=True)
        _write_slot_excel(sheet, g["docs"], g["grid"], g["grid"].get("cells", {}),
                          openpyxl_mod, provenance=provenance)
    return True


def _document_hint_of(ed):
    """The decimal separator a stored result's own numbers agree on, or None."""
    from amounts import document_hint
    if not isinstance(ed, dict):
        return None
    texts = list((ed.get("extracted_fields") or {}).values())
    for k, rows in ed.items():
        if k.endswith("_rows") and isinstance(rows, list):
            texts += [v for r in rows if isinstance(r, dict)
                      for c, v in r.items() if not str(c).startswith("_")]
    return document_hint(texts, ed.get("locale"))


def _write_slot_excel(ws, doc_results, sheet_data, cells_tpl, openpyxl_mod,
                      provenance=True, marks=True):
    """
    Writer for slot-directed extraction.

    There is no placement logic here, and that is the point: every value was
    requested for a specific address and is written to that address. The only
    computed thing is the row shift when a document has more table rows than
    the template reserved blank rows for.

    REVIEW MARKS (`marks`, on by default). An export must never present a
    value a person still owes a look — unverified, low, empty and flagged, or
    carrying `review_required` — as a clean final value. Each such cell gets a
    fill and a comment naming its status and reasons (app/core/review.py, the
    same classification the review screen uses); an edited or reviewed value
    gets its own fill. A "Review status" column right of the sheet says, per
    row, whether anything in it needs review, and its header counts the
    values that do. Returns that count.

    It also carries provenance (I2), unless `provenance=False`: a Source column
    right of the template naming the document and file page of every row that
    holds extracted values, and a comment on every extracted value quoting the
    words it was read from. "Values only" meant a merged block, a wrong region
    or a second document arrived with nothing saying where it came from, and
    the file is where an accountant checks a figure. A job stored before this
    carries no provenance and exports exactly as it did.
    """
    from openpyxl.cell import MergedCell

    # The sheet's extent is the static text PLUS every slot the shape found.
    # _find_template_dimensions counts only cells that carry text, which was
    # right when an empty value cell was flagged extractTarget. Under the one
    # rule (text = label, empty = slot) a value cell carries neither text nor a
    # flag, so a template whose value column has no heading — a plain list of
    # "Label | <empty>" rows — collapsed to column A and every extracted value
    # was dropped at export. The slots say where the values go, so they define
    # the extent alongside the labels.
    max_r, max_c = _find_template_dimensions(cells_tpl)
    # A merge is part of the extent too. `_find_template_dimensions` counts
    # content cells, and a heading merged across A:D has content only in A —
    # so the sheet stopped short of the range and the merge was dropped for
    # being outside it. `template_shape._used_range` already widens for merges
    # on the shape side; this is the same fact on the writer side.
    for key, span in (sheet_data.get("merges") or {}).items():
        try:
            mr, mc = map(int, str(key).split(","))
            max_r = max(max_r, mr + max(1, int((span or {}).get("rows") or 1)) - 1)
            max_c = max(max_c, mc + max(1, int((span or {}).get("cols") or 1)) - 1)
        except (ValueError, AttributeError, TypeError):
            continue
    for doc in doc_results:
        ed = doc.get_extracted_data()
        if not isinstance(ed, dict):
            continue
        sm = ed.get("slot_map") or {}
        for f in sm.get("fields") or []:
            rc = _ref_to_rowcol(f.get("ref", ""))
            if rc:
                max_r, max_c = max(max_r, rc[0]), max(max_c, rc[1])
        for t in sm.get("tables") or []:
            max_r = max(max_r, int(t.get("end_row", 0)))
            if t.get("orientation") == "columns":
                # A transposed table overflows SIDEWAYS: more records than the
                # template drew columns for widens the sheet instead of
                # lengthening it.
                records = len(ed.get(f"{t['name']}_rows") or [])
                max_c = max(max_c, int(t.get("start_col", 0)) + max(records - 1, 0),
                            int(t.get("end_col", 0)))
            for col in t.get("columns") or []:
                max_c = max(max_c, int(col.get("col", 0)))

    GAP_BETWEEN_DOCS = 2
    doc_offset = 0
    # I2 — the Source column sits one past the template's full extent, so it
    # can never land on a cell the template or a band owns.
    source_col = max_c + 1
    # The Review status column comes after it (or takes its place when the
    # export carries no provenance).
    status_col = source_col + 1 if provenance else source_col
    from app.core import review as _review
    row_status = {}           # output row -> [status word, ...]
    unresolved_total = 0

    #: the decimal separator the document being written uses, from its own
    #: unambiguous numbers (engine/amounts.document_hint); per document.
    doc_hint = [None]

    def put(r, c, value, label=""):
        cell = ws.cell(row=r + 1, column=c + 1)
        if isinstance(cell, MergedCell):
            return None
        out = coerce_cell_value(value, label, doc_hint[0])
        if out is not None:
            cell.value = out
            # D8 — the cell keeps the number AND the notation the document
            # printed it with. Without this the export drops the currency
            # symbol and the trailing cents, which is what an accountant
            # actually looks at.
            fmt = cell_format(value, label, doc_hint[0])
            if fmt:
                cell.number_format = fmt
            precision_flag(cell, value, doc_hint[0])
            return cell
        return None

    for doc in doc_results:
        ed = doc.get_extracted_data()
        doc_hint[0] = _document_hint_of(ed)
        if not isinstance(ed, dict):
            continue
        slot_map = ed.get("slot_map") or {}
        tables = slot_map.get("tables") or []
        fields = ed.get("extracted_fields") or {}

        filename = str(getattr(doc, "filename", "") or "")
        field_prov = ed.get("field_provenance") or {}
        conf_map = (ed.get("validation") or {}).get("confidence_map") or {}
        source_rows = {}          # output row -> set of pages
        doc_reviewed = _review.document_reviewed(ed)
        flagged_refs = {str(f.get("ref")) for f in
                        ((ed.get("validation") or {}).get("flagged_fields") or [])
                        if isinstance(f, dict)}

        def mark(cell, r, c, st):
            """The review mark on one written value (see the docstring)."""
            nonlocal unresolved_total
            if not marks:
                return
            if not st["resolved"]:
                unresolved_total += 1
            m = _review.export_mark(st)
            if r is not None:
                row_status.setdefault(r, []).append(m[0] if m else "OK")
            if m is None:
                return
            if cell is None:
                cell = ws.cell(row=r + 1, column=c + 1) if r is not None else None
                if cell is None or isinstance(cell, MergedCell):
                    return
            from openpyxl.comments import Comment
            from openpyxl.styles import PatternFill
            cell.fill = PatternFill(fill_type="solid", fgColor=m[1])
            if st["status"] == "edited" and cell.comment:
                return            # the edited note already says so, and what it replaced
            text = m[2] + (f"\n\n{cell.comment.text}" if cell.comment else "")
            w, h = provenance_note_size(text)
            cell.comment = Comment(text, "DocAgent", width=w, height=h)

        def attach(cell, source, page, grounded, confidence, out_r,
                   edited=False, original=""):
            """The comment on one written value, and its row's Source entry.

            An EDITED value was typed by a person. It gets a comment saying so
            and naming what it replaced, never the replaced words presented as
            its source, and its page does not go into the row's Source cell:
            the value did not come from that page.
            """
            if not provenance or cell is None or not filename:
                return
            from openpyxl.comments import Comment
            if edited:
                text = provenance_edited_note(original, page, filename)
            else:
                text = provenance_note(source, page, filename, grounded=grounded,
                                       confidence=confidence)
            text = keep_precision_note(cell, text)
            w, h = provenance_note_size(text)
            cell.comment = Comment(text, "DocAgent", width=w, height=h)
            if out_r is not None:
                source_rows.setdefault(out_r, set()).update(
                    {page} if page and not edited else set())

        # how far each template row moves down, given table overflow. A
        # transposed table never pushes rows down — it grows to the right — so
        # it contributes no shift.
        shifts = []
        for t in tables:
            if t.get("orientation") == "columns":
                continue
            rows = ed.get(f"{t['name']}_rows") or []
            band = t["end_row"] - t["start_row"] + 1
            shifts.append((t["end_row"], max(0, len(rows) - band)))

        def out_row(tr):
            return tr + sum(ov for end, ov in shifts if tr > end)

        # Cells the table itself will fill, which must not also get static text.
        # A row-oriented band owns whole rows; a transposed one owns only its
        # record columns — its heading column IS the template's static labels
        # and has to survive.
        band_rows, band_cells = set(), set()
        for t in tables:
            if t.get("orientation") == "columns":
                band_cells.update(
                    (r, c) for r in range(t["start_row"], t["end_row"] + 1)
                    for c in range(t["start_col"], max_c + 1))
            else:
                band_rows.update(range(t["start_row"], t["end_row"] + 1))

        # 1. static template text, at its shifted position
        for key, cell_def in cells_tpl.items():
            if not isinstance(cell_def, dict) or cell_def.get("mergeParent"):
                continue
            try:
                tr, tc = map(int, str(key).split(","))
            except ValueError:
                continue
            if tr > max_r or tc > max_c or tr in band_rows:
                continue
            if (tr, tc) in band_cells:
                continue
            text = str(cell_def.get("value") or "").strip()
            if text:
                put(doc_offset + out_row(tr), tc, text)

        # 1b. THE LOOK — including inside a band.
        #
        # Style used to be applied in the loop above, which `continue`s past
        # every band row before it gets there. A TABLE TEMPLATE BORDERS ITS
        # BODY: the band is the part of the sheet a user draws a box around,
        # and it was the one part that arrived unstyled. D13 was verified on a
        # form template, which has no band, so the gap was invisible.
        #
        # The look travels with a cell whether or not it holds text — a
        # bordered empty value cell is a box the user drew, and it is the cells
        # WITHOUT text that the slots fill.
        styles = {}
        for key, cell_def in cells_tpl.items():
            if not isinstance(cell_def, dict) or cell_def.get("mergeParent"):
                continue
            try:
                tr, tc = map(int, str(key).split(","))
            except ValueError:
                continue
            if cell_def.get("style"):
                styles[(tr, tc)] = cell_def["style"]

        def dress(out_r, out_c, src):
            st = styles.get(src)
            if not st:
                return
            cell = ws.cell(row=doc_offset + out_r + 1, column=out_c + 1)
            if not isinstance(cell, MergedCell):
                _apply_cell_style(cell, st, openpyxl_mod)

        for (tr, tc) in styles:
            if tr > max_r or tc > max_c:
                continue
            if tr in band_rows or (tr, tc) in band_cells:
                continue
            dress(out_row(tr), tc, (tr, tc))

        # A band grows to the document's row count, so the rows the user drew
        # are not the rows that get written. Each drawn row keeps its own look;
        # rows BEYOND what was drawn take the look of the band's FIRST body
        # row, because that is the row the user drew as "a row of this table".
        # Repeating the LAST one instead would stamp a closing bottom border
        # onto every overflow row and draw the table as a stack of one-row
        # boxes.
        for t in tables:
            rows_n = len(ed.get(f"{t['name']}_rows") or [])
            if t.get("orientation") == "columns":
                first, last = t["start_col"], t["end_col"]
                for i in range(max(last - first + 1, rows_n)):
                    src_c = first + i if i <= last - first else first
                    for r in range(t["start_row"], t["end_row"] + 1):
                        dress(out_row(r), first + i, (r, src_c))
                continue
            height = t["end_row"] - t["start_row"] + 1
            for i in range(max(height, rows_n)):
                src_r = t["start_row"] + (i if i < height else 0)
                for c in range(0, max_c + 1):
                    dress(out_row(t["start_row"]) + i, c, (src_r, c))

        # 2. field slots, each at the exact address it was asked for
        slot_label = {f.get("ref"): f.get("row_label", "")
                      for f in (slot_map.get("fields") or [])}
        for ref, value in fields.items():
            rc = _ref_to_rowcol(ref)
            if rc is None:
                continue
            tr, tc = rc
            if tr > max_r or tc > max_c:
                continue
            out_r = doc_offset + out_row(tr)
            cell = put(out_r, tc, value, slot_label.get(ref, ""))
            prov = field_prov.get(ref)
            if prov:
                attach(cell, prov.get("source"), prov.get("page"),
                       prov.get("grounded"), conf_map.get(ref), out_r,
                       edited=bool(prov.get("edited")) or conf_map.get(ref) == "edited",
                       original=prov.get("original_value", ""))
            p = prov or {}
            mark(cell, out_r, tc, _review.value_status(
                value, conf_map.get(ref), review_required=bool(p.get("review_required")),
                reasons=p.get("review_reasons") or [],
                flagged=slot_label.get(ref, "") in flagged_refs or ref in flagged_refs,
                reviewed=doc_reviewed or bool(p.get("reviewed"))))

        # 3. table rows, in their band, by column header address
        for t in tables:
            rows = ed.get(f"{t['name']}_rows") or []
            if t.get("orientation") == "columns":
                # Transposed: record i is a COLUMN, each field a fixed row. A
                # record has no row of its own for a Source cell, so its
                # provenance travels in the comments only.
                for i, rec in enumerate(rows):
                    c = t["start_col"] + i
                    for f in t.get("fields") or []:
                        cell = put(doc_offset + out_row(f["row"]), c,
                                   rec.get(f["header"], ""), f.get("header", ""))
                        if "_source" in rec:
                            attach(cell, rec.get("_source"), rec.get("_page"),
                                   f["header"] not in (rec.get("_ungrounded") or []),
                                   rec.get("_confidence"), None,
                                   edited=f["header"] in (rec.get("_edited") or []),
                                   original=(rec.get("_original") or {}).get(f["header"], ""))
                        if cell is not None:
                            mark(cell, None, c, _review.value_status(
                            rec.get(f["header"], ""),
                            "edited" if f["header"] in (rec.get("_edited") or []) else (
                                "low" if f["header"] in (rec.get("_ungrounded") or [])
                                else rec.get("_confidence")),
                            review_required=bool(rec.get("_review_required")),
                            reasons=rec.get("_review_reasons") or [],
                            reviewed=doc_reviewed or bool(rec.get("_reviewed"))))
                continue
            header_row = t.get("header_row", t["start_row"])
            if provenance and rows and header_row < t["start_row"] and any(
                    "_source" in r for r in rows):
                hc = ws.cell(row=doc_offset + out_row(header_row) + 1,
                             column=source_col + 1)
                if not isinstance(hc, MergedCell) and hc.value is None:
                    hc.value = "Source"
            for i, row in enumerate(rows):
                r = doc_offset + out_row(t["start_row"]) + i
                for col in t["columns"]:
                    key = col.get("key") or col["header"]
                    cell = put(r, col["col"], row.get(key, ""), col.get("header", ""))
                    if "_source" in row:
                        attach(cell, row.get("_source"), row.get("_page"),
                               key not in (row.get("_ungrounded") or []),
                               row.get("_confidence"), r,
                               edited=key in (row.get("_edited") or []),
                               original=(row.get("_original") or {}).get(key, ""))
                    st = _review.value_status(
                            row.get(key, ""),
                            "edited" if key in (row.get("_edited") or []) else (
                                "low" if key in (row.get("_ungrounded") or [])
                                else row.get("_confidence")),
                            review_required=bool(row.get("_review_required")),
                            reasons=row.get("_review_reasons") or [],
                            reviewed=doc_reviewed or bool(row.get("_reviewed")))
                    if cell is not None:
                        mark(cell, r, col["col"], st)

        # 3b. one Source cell per row that received a value with provenance
        for out_r, pages in source_rows.items():
            sc = ws.cell(row=out_r + 1, column=source_col + 1)
            if isinstance(sc, MergedCell) or sc.value is not None:
                continue
            shown = sorted(pages)
            sc.value = (f"{filename} · p.{', '.join(str(p) for p in shown)}"
                        if shown else filename)

        # 4. the template's own merges, last, so nothing is written into a
        #    cell that a merge has since swallowed
        def _addr(tr, tc):
            if tr > max_r or tc > max_c:
                return None
            return (doc_offset + out_row(tr) + 1, tc + 1)

        _apply_template_merges(ws, sheet_data.get("merges"), _addr, band_rows,
                               openpyxl_mod)

        block_height = max_r + 1 + sum(ov for _, ov in shifts)
        doc_offset += block_height + GAP_BETWEEN_DOCS

    # 5. the Review status column: a header counting what needs review, and
    #    one word per row that received a value
    if marks:
        head = (f"Review status — {unresolved_total} value"
                f"{'' if unresolved_total == 1 else 's'} need"
                f"{'s' if unresolved_total == 1 else ''} review"
                if unresolved_total else "Review status — nothing needs review")
        for r, words in sorted(row_status.items()):
            word = next((w for w in ("Needs review", "Edited", "Reviewed")
                         if w in words), "OK")
            if r == 0:
                head += f" (row 1: {word})"
                continue
            sc = ws.cell(row=r + 1, column=status_col + 1)
            if not isinstance(sc, MergedCell) and sc.value is None:
                sc.value = word
        hc = ws.cell(row=1, column=status_col + 1)
        if not isinstance(hc, MergedCell):
            hc.value = head
    return unresolved_total
