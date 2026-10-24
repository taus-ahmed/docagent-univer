"""Review marks on an exported sheet, for every writer.

The slot writer marks values inline (`_write_slot_excel`). Every other writer
— the legacy form, table, mixed and layout writers, the table-rows fallback,
the flat table and the combined / per-file sheets — writes through a
`SheetMarks`: at the moment a writer puts a document's value into a cell it
calls `field` or `row_cell`, and the cell gets the same fill and comment the
slot writer gives it. The status comes from `app.core.review` (the function
the review screen uses), so a sheet and the screen cannot disagree.

`finish` adds the "Review status" column one past the sheet's last column: a
header counting the values that still need review, and one word per row that
received a value.
"""
from __future__ import annotations

from typing import Any, Optional

from app.core import review as R


def note_size(text: str) -> tuple[int, int]:
    """(width, height) in points for a comment box that shows `text` whole —
    the same sizing as `extract.provenance_note_size`."""
    width = 360
    per_line = width // 6
    lines = sum(max(1, -(-len(ln) // per_line)) for ln in text.splitlines())
    return width, 13 * lines + 12


class SheetMarks:
    def __init__(self, ws, enabled: bool = True):
        self.ws = ws
        self.enabled = enabled
        self.rows: dict[int, list[str]] = {}   # 1-based row -> status words
        self.unresolved = 0

    def mark(self, cell, st: dict) -> None:
        """Fill and comment one written cell by its status."""
        from openpyxl.cell import MergedCell
        if not self.enabled or cell is None or isinstance(cell, MergedCell):
            return
        if not st["resolved"]:
            self.unresolved += 1
        m = R.export_mark(st)
        self.rows.setdefault(cell.row, []).append(m[0] if m else "OK")
        if m is None:
            return
        from openpyxl.comments import Comment
        from openpyxl.styles import PatternFill
        cell.fill = PatternFill(fill_type="solid", fgColor=m[1])
        text = m[2] + (f"\n\n{cell.comment.text}" if cell.comment else "")
        w, h = note_size(text)
        cell.comment = Comment(text, "DocAgent", width=w, height=h)

    def field(self, cell, ed: Any, ref: Optional[str] = None,
              label: Optional[str] = None, value: Any = R._MISSING) -> None:
        self.mark(cell, R.field_status(ed, ref=ref, label=label, value=value))

    def row_cell(self, cell, ed: Any, row: Any, col: str) -> None:
        self.mark(cell, R.row_cell_status(ed, row, col))

    def finish(self, header_row: int = 1) -> int:
        """Write the Review status column; returns the unresolved count."""
        from openpyxl.cell import MergedCell
        if not self.enabled or not self.rows:
            return self.unresolved
        col = (self.ws.max_column or 0) + 1
        n = self.unresolved
        head = (f"Review status — {n} value{'' if n == 1 else 's'} "
                f"need{'s' if n == 1 else ''} review" if n
                else "Review status — nothing needs review")
        for r, words in sorted(self.rows.items()):
            word = next((w for w in ("Needs review", "Edited", "Reviewed")
                         if w in words), "OK")
            if r == header_row:
                head += f" (row {r}: {word})"
                continue
            sc = self.ws.cell(row=r, column=col)
            if not isinstance(sc, MergedCell) and sc.value is None:
                sc.value = word
        hc = self.ws.cell(row=header_row, column=col)
        if not isinstance(hc, MergedCell):
            hc.value = head
        return n
